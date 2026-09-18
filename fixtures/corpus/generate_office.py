"""Dev-only fixture generator for DOCX/PPTX/XLSX test fixtures (see
SPEC.md §5.1 fixtures/corpus, §7 M2). Not a runtime dependency.

Usage: python fixtures/corpus/generate_office.py
Requires: pip install python-docx openpyxl python-pptx
"""

from pathlib import Path

import docx
import openpyxl
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).parent


def make_docx(path: Path) -> None:
    document = docx.Document()
    document.add_heading("Project Notes", level=1)
    document.add_paragraph("The migration to the new warehouse is on schedule.")
    document.add_paragraph("Segunda seccion: la reunion sera el martes en la manana.")
    document.save(str(path))


def make_pptx(path: Path) -> None:
    presentation = Presentation()
    slide_layout = presentation.slide_layouts[1]

    slide1 = presentation.slides.add_slide(slide_layout)
    slide1.shapes.title.text = "Quarterly Kickoff"
    slide1.placeholders[1].text = "Welcome to the quarterly planning session."

    slide2 = presentation.slides.add_slide(slide_layout)
    slide2.shapes.title.text = "Presupuesto"
    slide2.placeholders[1].text = "El presupuesto total es de cinco mil euros."

    presentation.save(str(path))


def make_xlsx(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Inventory"
    sheet.append(["Item", "Quantity", "Location"])
    sheet.append(["Widget", 42, "Warehouse A"])
    sheet.append(["Gadget", 7, "Warehouse B"])

    sheet2 = workbook.create_sheet("Facturas")
    sheet2.append(["Cliente", "Total"])
    sheet2.append(["Electricista Lopez", 150])

    workbook.save(str(path))


def main() -> None:

    office_dir = ROOT / "office"
    office_dir.mkdir(parents=True, exist_ok=True)

    office_dir = ROOT / "office"
    make_docx(office_dir / "notes.docx")
    make_pptx(office_dir / "kickoff.pptx")
    make_xlsx(office_dir / "inventory.xlsx")
    print("wrote:", office_dir / "notes.docx")
    print("wrote:", office_dir / "kickoff.pptx")
    print("wrote:", office_dir / "inventory.xlsx")


if __name__ == "__main__":
    main()
